"""User-entered text must remain inert when reports are opened as spreadsheets."""

import csv
import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.reports.export import MONEY, Column, Sheet, to_csv_bytes, to_xlsx

PAYLOADS = [
    "=1+1",
    "+1+1",
    "-1+1",
    "@SUM(1,1)",
    "\t=1+1",
    "\r=1+1",
    "\n=1+1",
    "  =1+1",
    "＝1+1",
    "＋1+1",
    "－1+1",
    "＠SUM(1,1)",
    '=HYPERLINK("https://example.com","sample")',
]


@pytest.mark.parametrize("text", PAYLOADS + ["#REF!"])
def test_xlsx_text_is_literal_in_headers_rows_and_totals(text):
    sheet = Sheet(
        columns=[Column(text), Column("Amount", MONEY)],
        rows=[[text, Decimal("-125.50")]],
        totals=[text, Decimal("-125.50")],
    )
    wb = load_workbook(io.BytesIO(to_xlsx({"Report": sheet})), data_only=False)
    ws = wb.active
    for address in ["A1", "A2", "A3"]:
        assert ws[address].value == text
        assert ws[address].data_type == "s"
    assert ws["B2"].data_type == "n"
    assert ws["B2"].value == -125.5


@pytest.mark.parametrize("text", PAYLOADS)
def test_csv_quotes_and_neutralizes_formula_text_without_altering_numbers(text):
    data = to_csv_bytes([[text, Decimal("-125.50"), "Ordinary text", None]])
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows == [["\t" + text, "-125.50", "Ordinary text", ""]]
    assert data.decode("utf-8-sig").startswith('"\t')


def test_legacy_xlsx_list_rows_are_also_literal():
    wb = load_workbook(io.BytesIO(to_xlsx({"Report": [["=1+1"], ["=1+1"]]})))
    assert wb.active["A1"].data_type == "s"
    assert wb.active["A2"].data_type == "s"


@pytest.mark.django_db
@pytest.mark.parametrize("staff", [True, False])
def test_admin_inventory_exports_neutralize_stored_product_text(staff):
    from apps.inventory.models import Category, Product, ProductVariant
    from apps.inventory.resources import ProductVariantResource, StaffProductVariantResource

    category = Category.objects.create(name="Example category")
    product = Product.objects.create(category=category, name="=1+1")
    variant = ProductVariant.objects.create(
        product=product, sku="DEMO-SECURITY", cost_price=Decimal("100"), sale_price=Decimal("150")
    )
    resource = StaffProductVariantResource() if staff else ProductVariantResource()
    data = resource.export(ProductVariant.objects.filter(pk=variant.pk))
    rows = list(csv.reader(io.StringIO(data.csv)))
    index = rows[0].index("product__name")
    assert rows[1][index] == "\t=1+1"
    wb = load_workbook(io.BytesIO(data.xlsx), data_only=False)
    cell = wb.active.cell(row=2, column=index + 1)
    assert cell.value == "\t=1+1"
    assert cell.data_type == "s"
    assert ("cost_price" in rows[0]) is not staff
